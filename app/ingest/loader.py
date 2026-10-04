"""多格式文档解析：TXT / MD / PDF / Word / HTML / 网页链接。"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

from app.ingest.url_guard import UnsafeUrlError, fetch_url
from app.text import md5_hex

logger = logging.getLogger(__name__)

_DECODE_ORDER = ("utf-8-sig", "utf-8", "gb18030", "big5", "utf-16", "latin-1")
_STRIP_TAGS = ["script", "style", "noscript", "iframe", "svg", "canvas", "form", "nav", "footer", "aside"]


class UnsupportedFormatError(Exception):
    """不支持的文件格式 / 解析失败。"""


@dataclass(slots=True)
class LoadedDocument:
    """解析结果。"""

    name: str
    text: str
    source_type: str = "file"
    source_uri: str | None = None
    meta: dict = field(default_factory=dict)
    content_hash: str = ""

    def finalize(self) -> "LoadedDocument":
        if not self.content_hash:
            self.content_hash = md5_hex(self.text.encode("utf-8", errors="ignore"))
        return self


def decode_bytes(data: bytes) -> str:
    for encoding in _DECODE_ORDER:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="ignore")


def load_from_path(path: str | Path, *, name: str | None = None) -> LoadedDocument:
    file_path = Path(path)
    if not file_path.exists():
        raise UnsupportedFormatError(f"文件不存在：{file_path}")
    suffix = file_path.suffix.lower()
    display_name = name or file_path.name

    if suffix in {".txt", ".md", ".markdown", ".text", ".log", ".csv"}:
        text = decode_bytes(file_path.read_bytes())
        return LoadedDocument(display_name, text, "file", str(file_path), {"suffix": suffix}).finalize()
    if suffix == ".pdf":
        return _load_pdf(file_path.read_bytes(), display_name, str(file_path))
    if suffix == ".docx":
        return _load_docx(file_path, display_name)
    if suffix == ".doc":
        raise UnsupportedFormatError(
            "暂不支持 .doc 旧版二进制格式，请另存为 .docx 后再上传（或用 Word/WPS 另存为）"
        )
    if suffix in {".html", ".htm"}:
        return _load_html(file_path.read_bytes(), display_name, str(file_path))
    raise UnsupportedFormatError(
        f"不支持的文件类型：{suffix or '未知'}，当前支持 TXT / MD / PDF / DOCX / HTML 以及网页链接"
    )


def load_from_text(text: str, name: str = "手动输入文本") -> LoadedDocument:
    return LoadedDocument(name, text, "text", None, {"suffix": ".txt"}).finalize()


def load_from_url(
    url: str,
    *,
    timeout: float = 30.0,
    name: str | None = None,
    max_bytes: int = 10 * 1024 * 1024,
    max_redirects: int = 3,
    allow_private: bool = False,
) -> LoadedDocument:
    """抓取网页并解析。所有请求都经过 url_guard 的安全校验与体积限制。"""

    try:
        result = fetch_url(
            url,
            timeout=timeout,
            max_bytes=max_bytes,
            max_redirects=max_redirects,
            allow_private=allow_private,
        )
    except UnsafeUrlError as exc:
        raise UnsupportedFormatError(f"网页抓取被拒绝：{exc}") from exc
    except Exception as exc:  # noqa: BLE001 - 统一转成可读错误
        raise UnsupportedFormatError(f"网页抓取失败：{exc}") from exc

    content_type = result.content_type
    final_url = result.url
    if "application/pdf" in content_type or final_url.lower().endswith(".pdf"):
        return _load_pdf(result.content, name or final_url, final_url)
    if "text/html" in content_type or "text/plain" not in content_type:
        return _load_html(result.content, name or final_url, final_url)
    text = decode_bytes(result.content)
    return LoadedDocument(name or final_url, text, "url", final_url, {"content_type": content_type}).finalize()


# --------------------------------------------------------------------------- #
def _load_pdf(data: bytes, name: str, uri: str) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError("未安装 pypdf，无法解析 PDF：pip install pypdf") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                raise UnsupportedFormatError("PDF 已加密，无法解析") from None
        pages: list[str] = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:  # noqa: BLE001 - 单页失败不影响整体
                logger.warning("PDF 某页解析失败：%s", exc)
                pages.append("")
    except UnsupportedFormatError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedFormatError(f"PDF 解析失败：{exc}") from exc

    text = "\n\n".join(page.strip() for page in pages if page.strip())
    if not text.strip():
        raise UnsupportedFormatError("PDF 未提取到文本（可能是扫描件，需要 OCR 后再上传）")
    metadata = {}
    try:
        info = reader.metadata or {}
        metadata = {"title": getattr(info, "title", None), "author": getattr(info, "author", None)}
    except Exception:  # noqa: BLE001
        pass
    title = (metadata.get("title") or "").strip() or name
    return LoadedDocument(
        title, text, "file", uri, {"suffix": ".pdf", "pages": len(reader.pages), **metadata}
    ).finalize()


def _load_docx(path: Path, name: str) -> LoadedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError("未安装 python-docx，无法解析 Word：pip install python-docx") from exc

    try:
        document = docx.Document(str(path))
        parts: list[str] = []
        for kind, block in _iter_docx_blocks(document):
            if kind == "table":
                for row in block.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        parts.append("| " + " | ".join(cells) + " |")
                parts.append("")
            else:
                text = block.text.strip()
                if not text:
                    parts.append("")
                    continue
                style = (block.style.name or "").lower() if block.style is not None else ""
                parts.append(_docx_heading_prefix(style) + text if style.startswith("heading") else text)
            parts.append("")
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedFormatError(f"Word 解析失败：{exc}") from exc

    text = "\n".join(parts)
    return LoadedDocument(name, text, "file", str(path), {"suffix": ".docx"}).finalize()


def _docx_heading_prefix(style_name: str) -> str:
    digits = "".join(ch for ch in style_name if ch.isdigit())
    level = int(digits) if digits.isdigit() and 1 <= int(digits) <= 6 else 1
    return "#" * level + " "


def _iter_docx_blocks(document):
    """按文档真实顺序同时产出段落和表格。"""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, document)


def _load_html(data: bytes, name: str, uri: str) -> LoadedDocument:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormatError("未安装 beautifulsoup4，无法解析网页") from exc

    html = decode_bytes(data)
    try:
        soup = BeautifulSoup(html, "lxml")
    except Exception:  # noqa: BLE001 - 没有 lxml 时退回标准库解析器
        soup = BeautifulSoup(html, "html.parser")

    for tag in soup(_STRIP_TAGS):
        tag.decompose()

    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    elif soup.h1:
        title = soup.h1.get_text(" ", strip=True)

    # 表格 → markdown 行，保证切片器能识别为表格块
    for row in soup.find_all("tr"):
        cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["td", "th"])]
        if cells:
            row.replace_with(f"\n| {' | '.join(cells)} |\n")

    for level in range(1, 7):
        for tag in soup.find_all(f"h{level}"):
            tag.replace_with(f"\n{'#' * level} {tag.get_text(' ', strip=True)}\n")

    for tag in soup.find_all("li"):
        tag.insert_before("\n- ")
    for tag in soup.find_all(["p", "div", "section", "article", "br"]):
        tag.insert_after("\n")

    text = soup.get_text("\n")
    lines = [line.strip() for line in text.split("\n")]
    cleaned: list[str] = []
    for line in lines:
        if line or (cleaned and cleaned[-1]):
            cleaned.append(line)
    text = "\n".join(cleaned).strip()
    if not text:
        raise UnsupportedFormatError("网页未提取到正文内容")
    # 文件名形如 xxx.html 时，优先用网页标题作为文档名，检索结果更易读
    display_name = name if name and not name.lower().endswith((".html", ".htm")) else (title or name or uri)
    return LoadedDocument(
        display_name, text, "url", uri, {"suffix": ".html", "title": title}
    ).finalize()
